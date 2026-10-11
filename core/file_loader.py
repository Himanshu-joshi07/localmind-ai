"""Turn an uploaded file into plain text, enforcing the user's size and page limits."""

from dataclasses import dataclass
from io import BytesIO

from config import HARD_UPLOAD_LIMIT_MB, SUPPORTED_FILE_TYPES

WORDS_PER_PAGE = 500
CSV_PREVIEW_ROWS = 200


class FileLoadError(Exception):
    """A problem the user can fix; the message is shown as-is in the UI."""


@dataclass
class LoadedFile:
    name: str
    kind: str
    text: str
    pages: int | None
    size_mb: float


def load_file(name: str, data: bytes, max_file_mb: int, max_pages: int) -> LoadedFile:
    size_mb = len(data) / (1024 * 1024)
    if size_mb > max_file_mb:
        raise FileLoadError(
            f"This file is {size_mb:.1f} MB, but your limit is {max_file_mb} MB. "
            f"You can raise it in 📎 File settings (up to {HARD_UPLOAD_LIMIT_MB} MB)."
        )

    kind = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if kind not in SUPPORTED_FILE_TYPES:
        raise FileLoadError(f"`.{kind}` files aren't supported. Use: {', '.join(SUPPORTED_FILE_TYPES)}.")

    try:
        if kind == "pdf":
            text, pages = _read_pdf(data, max_pages)
        elif kind == "docx":
            text, pages = _read_docx(data, max_pages)
        elif kind == "csv":
            text, pages = _read_csv(data), None
        else:
            text, pages = data.decode("utf-8", errors="replace"), None
    except FileLoadError:
        raise
    except Exception as exc:
        raise FileLoadError(f"Couldn't read `{name}`: {exc}") from exc

    if not text.strip():
        hint = " It may be a scanned PDF (images of text), which needs OCR and isn't supported yet." if kind == "pdf" else ""
        raise FileLoadError(f"No readable text found in `{name}`.{hint}")

    return LoadedFile(name=name, kind=kind, text=text, pages=pages, size_mb=size_mb)


def _page_limit_error(pages: int, max_pages: int, approx: bool = False) -> FileLoadError:
    about = "about " if approx else ""
    return FileLoadError(
        f"This document has {about}{pages} page{'s' if pages != 1 else ''}, but your limit is {max_pages}. "
        f"You can raise it in 📎 File settings."
    )


def _read_pdf(data: bytes, max_pages: int) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = len(reader.pages)
    if pages > max_pages:
        raise _page_limit_error(pages, max_pages)
    text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    return text, pages


def _read_docx(data: bytes, max_pages: int) -> tuple[str, int]:
    from docx import Document

    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(parts)

    # Word files don't store a page count, so estimate it from the word count.
    pages = max(1, len(text.split()) // WORDS_PER_PAGE)
    if pages > max_pages:
        raise _page_limit_error(pages, max_pages, approx=True)
    return text, pages


def _read_csv(data: bytes) -> str:
    import pandas as pd

    df = pd.read_csv(BytesIO(data))
    rows, cols = df.shape
    shown = df.head(CSV_PREVIEW_ROWS)
    note = "" if rows <= CSV_PREVIEW_ROWS else f" (first {CSV_PREVIEW_ROWS} of {rows} rows shown)"
    return (
        f"CSV file with {rows} rows and {cols} columns.\n"
        f"Columns: {', '.join(map(str, df.columns))}\n\n"
        f"Summary statistics:\n{df.describe(include='all').to_string()}\n\n"
        f"Data{note}:\n{shown.to_csv(index=False)}"
    )
